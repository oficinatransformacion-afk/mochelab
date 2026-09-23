import argparse
import json
import re
import unicodedata
from pathlib import Path

import openpyxl


def normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


parser = argparse.ArgumentParser(description="Valida y prepara respuestas ATF para la carga local de Mochelab.")
parser.add_argument("--technical", required=True)
parser.add_argument("--soft", required=True)
parser.add_argument("--model", required=True)
parser.add_argument("--output", required=True)
args = parser.parse_args()

with open(args.model, encoding="utf-8") as source:
    model_file = json.load(source)
model = next(role for role in model_file["roles"] if role["sourceId"] == "MAT-07")
technical_section = next(section for section in model["sections"] if section["type"] == "TECHNICAL")
soft_section = next(section for section in model["sections"] if section["type"] == "SOFT")

technical_items = {}
for dimension in technical_section["dimensions"]:
    for item in dimension["items"]:
        technical_items[(normalize(dimension["name"]), normalize(item["statement"]))] = {
            "sectionCode": technical_section["code"], "dimensionCode": dimension["code"], "itemCode": item["code"]
        }

people = {}
technical_book = openpyxl.load_workbook(args.technical, read_only=True, data_only=True)
for dni, name, dimension, _level, statement, complies in technical_book.active.iter_rows(min_row=2, values_only=True):
    if dni is None:
        continue
    key = str(int(dni)) if isinstance(dni, float) and dni.is_integer() else str(dni).strip()
    statement_key = normalize(str(statement or "").lstrip("* "))
    item = technical_items.get((normalize(dimension), statement_key))
    if not item:
        raise ValueError(f"No se encontró el comportamiento técnico: {dimension} / {statement}")
    value = {"si": 2, "no": 0}.get(normalize(complies))
    if value is None:
        raise ValueError(f"Respuesta técnica inválida para DNI {key}: {complies}")
    person = people.setdefault(key, {"dni": key, "name": str(name).strip(), "email": None, "answers": []})
    person["answers"].append({**item, "value": value, "source": "ATF.xlsx"})

soft_dimensions = {normalize(dimension["name"]): dimension for dimension in soft_section["dimensions"]}
soft_book = openpyxl.load_workbook(args.soft, read_only=True, data_only=True)
rows = soft_book.active.iter_rows(values_only=True)
headers = next(rows)
for row in rows:
    dni, name, email, *values = row
    if dni is None:
        continue
    key = str(int(dni)) if isinstance(dni, float) and dni.is_integer() else str(dni).strip()
    if key not in people:
        raise ValueError(f"El DNI {key} aparece en ATF_BLANDO.xlsx pero no en ATF.xlsx")
    people[key]["email"] = str(email).strip() if email else None
    for header, value in zip(headers[3:], values):
        dimension = soft_dimensions.get(normalize(header))
        if not dimension:
            raise ValueError(f"Competencia blanda desconocida: {header}")
        if value not in (0, 1, 2):
            raise ValueError(f"Valor blando inválido para DNI {key}, {header}: {value}")
        item = dimension["items"][0]
        people[key]["answers"].append({
            "sectionCode": soft_section["code"], "dimensionCode": dimension["code"],
            "itemCode": item["code"], "value": int(value), "source": "ATF_BLANDO.xlsx",
        })

if len(people) != 6:
    raise ValueError(f"Se esperaban 6 personas y se encontraron {len(people)}")
for person in people.values():
    if len(person["answers"]) != 125:
        raise ValueError(f"El DNI {person['dni']} tiene {len(person['answers'])} respuestas; se esperaban 125")
    keys = {(answer["sectionCode"], answer["dimensionCode"], answer["itemCode"]) for answer in person["answers"]}
    if len(keys) != 125:
        raise ValueError(f"El DNI {person['dni']} contiene respuestas duplicadas")

output = Path(args.output)
output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(json.dumps({"modelVersion": model_file["version"], "people": list(people.values())}, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Archivo validado: {len(people)} personas y {sum(len(person['answers']) for person in people.values())} respuestas.")
